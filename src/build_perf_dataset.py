"""Build a length-stratified dataset for embedding THROUGHPUT benchmarking.

Reads the benchmark workbooks (and the corpus exports they came from) and emits
samples bucketed by real token count, so a run can be timed separately on short,
medium and long inputs. It measures speed, not accuracy: nothing here carries a
label or a relevance judgement.

    python src/build_perf_dataset.py
    python src/build_perf_dataset.py --per-group 500 --out-dir results/perf
    python src/build_perf_dataset.py --dry-run          # report availability only

Everything configurable sits in the CONFIG block below; the command line just
overrides a few of those.

Where each group comes from, and why:

  short   questions and single extracted facts - genuinely short real inputs.
  medium  one chunk of a document, the unit the RAG index actually stores.
  long    consecutive chunks of ONE document, in document order, concatenated until
          the total lands inside the target band. No document in this corpus reaches
          7000 tokens in a single chunk, and only 84 reach it in total, so a long
          sample is a real document prefix rather than a single chunk. Nothing is
          invented and nothing is stitched across documents.

Texts are NFC-normalised and de-duplicated exactly, then sampled evenly across
length sub-bins and across sources, so a group is not all one length or all one
kind of text. If a group cannot be filled, the script takes what exists and says
so - it never pads by repeating or inventing text.
"""
import argparse
import csv
import glob
import hashlib
import json
import os
import random
import sys
import unicodedata
from collections import defaultdict

# ----------------------------------------------------------------- CONFIG
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

# Inputs. Globs and paths are resolved relative to the project root.
XLSX_GLOB = "data/*.xlsx"
CORPUS_JSONL = ["data/corpus_full.jsonl", "data/corpus.jsonl"]

OUT_DIR = "results/perf_dataset"
OUT_BASENAME = "perf_dataset"
WRITE_PER_GROUP_FILES = True          # also one CSV per length_group

# Tokenizer used to define the groups. Qwen spends the most tokens on Vietnamese
# of the models under test, so bucketing with it keeps every other model at or
# below the stated band.
TOKENIZER = "Qwen/Qwen3-Embedding-0.6B"
# Extra tokenizers only add reference columns, e.g. "AITeamVN/Vietnamese_Embedding".
EXTRA_TOKENIZERS = {}                  # {"column_suffix": "hf_id"}

# Half-open token bands [lo, hi). Texts outside every band are dropped.
GROUPS = {
    "short": (1, 40),
    "medium": (500, 2000),
    "long": (7000, 8000),
}
SAMPLES_PER_GROUP = 200
LENGTH_BINS = 8                        # sub-bins per group, sampled round-robin
SEED = 20260916

# Column headers worth reading out of the workbooks, matched case-insensitively.
QUESTION_COLUMNS = ["cau_hoi", "cau_hoi_dai", "question"]
ANSWER_COLUMNS = ["cau_tra_loi", "expected", "claim"]
# Never read these: derived numbers, ids, truncated previews, stored vectors.
SKIP_COLUMN_SUBSTRINGS = ["text_300", "phan_", "chunk_id", "fact_id", "score",
                          "evidence", "reason"]

FACT_XLSX = "data/chunk_va_fact_500_bai.xlsx"
CHUNK_SHEET = "Chunk và fact"

MIN_CHARS = 8                          # ignore stubs and separator junk
TOKENIZE_BATCH = 256
# ------------------------------------------------------------- END CONFIG

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))


def abspath(p):
    return p if os.path.isabs(p) else os.path.join(ROOT, p)


def nfc(t):
    return unicodedata.normalize("NFC", str(t)).strip()


def looks_like_text(s):
    """Reject vectors, ids and separator rubbish that share a column with prose."""
    if not s or len(s) < MIN_CHARS:
        return False
    if s[0] in "[{" and s.rstrip().endswith(("]", "}")):
        return False                              # a stored JSON vector or blob
    letters = sum(ch.isalpha() for ch in s)
    return letters >= max(MIN_CHARS // 2, 0.3 * len(s))


class Counter:
    """Token counts from a fast tokenizer, batched."""

    def __init__(self, hf_id):
        from transformers import AutoTokenizer
        self.tk = AutoTokenizer.from_pretrained(hf_id)
        self.hf_id = hf_id

    def count(self, texts):
        out = []
        for i in range(0, len(texts), TOKENIZE_BATCH):
            enc = self.tk(texts[i:i + TOKENIZE_BATCH], add_special_tokens=False)
            out += [len(x) for x in enc["input_ids"]]
        return out


# ------------------------------------------------------------- collectors

def read_workbook_texts():
    """(text, source) from every workbook: questions, answers, facts, chunk text."""
    import openpyxl

    found = []
    for path in sorted(glob.glob(abspath(XLSX_GLOB))):
        base = os.path.basename(path)
        if base.startswith("~$") or ".bak." in base:
            continue                              # Excel lock file / our own backup
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        for sheet in wb.sheetnames:
            rows = wb[sheet].iter_rows(values_only=True)
            header = next(rows, None)
            if not header:
                continue
            hdr = [str(c).strip().lower() if c is not None else "" for c in header]
            wanted = {}
            for j, h in enumerate(hdr):
                if not h or any(s in h for s in SKIP_COLUMN_SUBSTRINGS):
                    continue
                if h in QUESTION_COLUMNS:
                    wanted[j] = "question"
                elif h in ANSWER_COLUMNS:
                    wanted[j] = "answer"
                elif "c.text" in h:
                    wanted[j] = "chunk"
            for row in rows:
                for j, kind in wanted.items():
                    if j < len(row) and isinstance(row[j], str):
                        t = nfc(row[j])
                        if looks_like_text(t):
                            found.append((t, kind))
        wb.close()
    return found


def read_facts():
    """Single fact sentences, and all of one chunk's facts joined.

    The joined form is how the workbook already groups them (one row per chunk); the
    per-fact metadata lines are dropped, so both forms are original sentences. The
    joined form gives the medium band a second kind of text alongside chunk prose.
    """
    from fact_xlsx import load_facts
    facts = load_facts(abspath(FACT_XLSX)) or []
    out = [(f["text"], "fact") for f in facts if looks_like_text(f["text"])]

    by_chunk = defaultdict(list)
    for f in facts:
        by_chunk[f["chunk_id"]].append((f["k"], f["text"]))
    for parts in by_chunk.values():
        joined = "\n".join(t for _, t in sorted(parts))
        if looks_like_text(joined):
            out.append((joined, "fact_block"))
    return out


def read_corpus_chunks():
    """(text, source) per chunk, plus {doc_id: [(chunk_index, text)]} for concatenation."""
    texts, by_doc, seen = [], defaultdict(dict), set()
    for rel in CORPUS_JSONL:
        path = abspath(rel)
        if not os.path.exists(path):
            continue
        with open(path, encoding="utf-8") as f:
            for line in f:
                r = json.loads(line)
                t = nfc(r.get("text") or "")
                if not looks_like_text(t):
                    continue
                if r["chunk_id"] not in seen:     # corpus.jsonl is a subset of the full file
                    seen.add(r["chunk_id"])
                    texts.append((t, "chunk"))
                by_doc[r["doc_id"]][r["chunk_index"]] = t
    return texts, by_doc


def build_document_prefixes(by_doc, counter, lo, hi):
    """Consecutive whole chunks of one document until the total lands in [lo, hi)."""
    samples = []
    for doc_id, chunks in by_doc.items():
        parts = [t for _, t in sorted(chunks.items())]
        counts = counter.count(parts)
        if sum(counts) < lo:
            continue                              # document is simply not long enough
        total, taken = 0, 0
        for n in counts:
            if total + n >= hi:
                break                             # the next chunk would overshoot
            total += n
            taken += 1
        if lo <= total < hi and taken:
            samples.append({"text": "\n\n".join(parts[:taken]), "source": "doc_prefix",
                            "doc_id": doc_id, "n_chunks": taken})
    return samples


# --------------------------------------------------------------- sampling

def stratified_pick(items, want, bins, rng):
    """Spread the pick over length sub-bins and over sources, without replacement."""
    if len(items) <= want:
        return list(items)
    lo = min(i["token_count"] for i in items)
    hi = max(i["token_count"] for i in items) + 1
    width = max((hi - lo) / bins, 1e-9)

    buckets = defaultdict(lambda: defaultdict(list))
    for it in items:
        b = min(int((it["token_count"] - lo) / width), bins - 1)
        buckets[b][it["source"]].append(it)
    for by_source in buckets.values():
        for lst in by_source.values():
            rng.shuffle(lst)

    picked, order = [], sorted(buckets)
    while len(picked) < want:
        progressed = False
        for b in order:                           # one pass: a turn per bin ...
            for src in sorted(buckets[b]):        # ... and per source inside it
                if buckets[b][src] and len(picked) < want:
                    picked.append(buckets[b][src].pop())
                    progressed = True
        if not progressed:
            break
    return picked


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--per-group", type=int, default=SAMPLES_PER_GROUP)
    ap.add_argument("--out-dir", default=OUT_DIR)
    ap.add_argument("--tokenizer", default=TOKENIZER)
    ap.add_argument("--seed", type=int, default=SEED)
    ap.add_argument("--dry-run", action="store_true", help="report availability, write nothing")
    args = ap.parse_args()

    rng = random.Random(args.seed)
    counter = Counter(args.tokenizer)
    print(f"tokenizer: {args.tokenizer}")

    # ---- gather candidates --------------------------------------------------
    pool = read_workbook_texts() + read_facts()
    corpus_texts, by_doc = read_corpus_chunks()
    pool += corpus_texts
    print(f"raw candidates: {len(pool)} from {len(set(s for _, s in pool))} kinds, "
          f"{len(by_doc)} documents")

    uniq = {}
    for text, source in pool:
        uniq.setdefault(hashlib.sha256(text.encode("utf-8")).hexdigest(),
                        {"text": text, "source": source})
    items = list(uniq.values())
    print(f"unique texts: {len(items)}")

    counts = counter.count([i["text"] for i in items])
    for it, n in zip(items, counts):
        it["token_count"] = n

    long_lo, long_hi = GROUPS["long"]
    doc_samples = build_document_prefixes(by_doc, counter, long_lo, long_hi)
    for s in doc_samples:
        s["token_count"] = counter.count([s["text"]])[0]
    doc_samples = [s for s in doc_samples if long_lo <= s["token_count"] < long_hi]
    print(f"document prefixes in [{long_lo}, {long_hi}): {len(doc_samples)}")

    # ---- bucket and sample --------------------------------------------------
    by_group, chosen = {}, {}
    for group, (lo, hi) in GROUPS.items():
        pool_g = [i for i in items if lo <= i["token_count"] < hi]
        if group == "long":
            pool_g = pool_g + doc_samples         # single chunks rarely reach this band
        by_group[group] = pool_g
        chosen[group] = stratified_pick(pool_g, args.per_group, LENGTH_BINS, rng)

    print(f"\n{'group':<8}{'band':>14}{'available':>11}{'picked':>8}   "
          f"{'tokens min/p50/max':>22}   sources")
    summary = {}
    for group, (lo, hi) in GROUPS.items():
        picked = sorted(chosen[group], key=lambda x: x["token_count"])
        ns = [p["token_count"] for p in picked]
        src = defaultdict(int)
        for p in picked:
            src[p["source"]] += 1
        span = (f"{ns[0]}/{ns[len(ns)//2]}/{ns[-1]}" if ns else "-")
        print(f"{group:<8}{f'[{lo},{hi})':>14}{len(by_group[group]):>11}{len(picked):>8}   "
              f"{span:>22}   {dict(src)}")
        if len(picked) < args.per_group:
            print(f"         only {len(picked)} of {args.per_group} available - "
                  f"taken as is, nothing invented to fill the gap")
        summary[group] = {"band": [lo, hi], "available": len(by_group[group]),
                          "picked": len(picked), "sources": dict(src),
                          "tokens": {"min": ns[0], "max": ns[-1]} if ns else None}

    if args.dry_run:
        print("\ndry run - nothing written")
        return

    # ---- extra tokenizer columns -------------------------------------------
    extra = {name: Counter(hf) for name, hf in EXTRA_TOKENIZERS.items()}

    rows = []
    for group in GROUPS:
        for i, p in enumerate(sorted(chosen[group], key=lambda x: x["token_count"]), 1):
            row = {"id": f"{group}-{i:04d}", "text": p["text"],
                   "token_count": p["token_count"], "length_group": group,
                   "source": p["source"], "char_count": len(p["text"]),
                   "doc_id": p.get("doc_id", ""), "n_chunks": p.get("n_chunks", 1),
                   "sha256": hashlib.sha256(p["text"].encode("utf-8")).hexdigest()[:16]}
            for name, c in extra.items():
                row[f"token_count_{name}"] = c.count([p["text"]])[0]
            rows.append(row)

    out_dir = abspath(args.out_dir)
    os.makedirs(out_dir, exist_ok=True)
    fields = list(rows[0])

    def write_csv(path, subset):
        with open(path, "w", encoding="utf-8-sig", newline="") as f:
            w = csv.DictWriter(f, fieldnames=fields)
            w.writeheader()
            w.writerows(subset)
        print(f"  {os.path.relpath(path, ROOT)}  ({len(subset)} rows)")

    print()
    write_csv(os.path.join(out_dir, f"{OUT_BASENAME}.csv"), rows)
    if WRITE_PER_GROUP_FILES:
        for group in GROUPS:
            write_csv(os.path.join(out_dir, f"{OUT_BASENAME}_{group}.csv"),
                      [r for r in rows if r["length_group"] == group])

    meta = {"tokenizer": args.tokenizer, "seed": args.seed,
            "samples_per_group": args.per_group, "groups": summary,
            "inputs": {"xlsx": XLSX_GLOB, "corpus": CORPUS_JSONL},
            "note": "throughput benchmark only - no labels, no relevance judgements"}
    with open(os.path.join(out_dir, f"{OUT_BASENAME}_summary.json"), "w",
              encoding="utf-8") as f:
        json.dump(meta, f, ensure_ascii=False, indent=2)
    print(f"  {os.path.relpath(os.path.join(out_dir, OUT_BASENAME + '_summary.json'), ROOT)}")


if __name__ == "__main__":
    main()
